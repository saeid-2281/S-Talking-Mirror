from __future__ import annotations

import csv
import io
import uuid
from collections import Counter
from pathlib import Path
from typing import Iterable

from app.csv_loader import ENCODINGS, FILENAME_EXT, INVALID_WINDOWS
from app.exceptions import CSVValidationError
from app.models.domain import TTSJob
from app.models.project_source import (
    ProjectSource,
    SourceCollectionImportResult,
    SourceColumnMapping,
    SourceImportIssue,
    SourceImportResult,
    SourceImportStatus,
    SourceType,
    source_hash,
    utc_now,
)

SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xlsm"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


class SourceImportService:
    """Imports validated project sources without mutating source files."""

    def create_sources(
        self,
        paths: Iterable[Path],
        *,
        project_id: int | None = None,
        starting_order: int = 0,
    ) -> list[ProjectSource]:
        sources: list[ProjectSource] = []
        order = starting_order
        for raw_path in paths:
            path = Path(raw_path)
            suffix = path.suffix.lower()
            if suffix == ".xls":
                sources.append(self._unsupported_source(path, project_id, order, "Legacy .xls requires conversion to .xlsx."))
                order += 1
                continue
            if suffix not in SUPPORTED_EXTENSIONS:
                sources.append(self._unsupported_source(path, project_id, order, f"Unsupported source type: {suffix or 'none'}."))
                order += 1
                continue
            if suffix in EXCEL_EXTENSIONS:
                worksheets = self.excel_worksheets(path)
                for worksheet_order, worksheet in enumerate(worksheets):
                    sources.append(
                        self._source(
                            path,
                            project_id=project_id,
                            order=order + worksheet_order,
                            source_type=SourceType(suffix.strip(".")),
                            worksheet_name=worksheet,
                        )
                    )
                order += max(len(worksheets), 1)
            else:
                sources.append(
                    self._source(
                        path,
                        project_id=project_id,
                        order=order,
                        source_type=SourceType.TSV if suffix == ".tsv" else SourceType.CSV,
                    )
                )
                order += 1
        return sources

    def import_sources(self, sources: list[ProjectSource]) -> SourceCollectionImportResult:
        results = [self.import_source(source) for source in sorted(sources, key=lambda item: item.import_order)]
        jobs: list[TTSJob] = []
        for result in results:
            if result.source.enabled:
                jobs.extend(result.jobs)
        collisions = self.detect_collisions(jobs)
        return SourceCollectionImportResult(results, jobs, collisions)

    def import_source(self, source: ProjectSource) -> SourceImportResult:
        if source.import_status == SourceImportStatus.UNSUPPORTED:
            return SourceImportResult(source, issues=[self._issue(source, "error", "unsupported_source", "Unsupported source type.")])
        if not source.source_path.exists():
            source.import_status = SourceImportStatus.MISSING
            return SourceImportResult(source, issues=[self._issue(source, "error", "missing_source", "Source file is missing.")])
        if source.source_type in {SourceType.CSV, SourceType.TSV}:
            return self._import_delimited(source)
        return self._import_excel(source)

    def detect_collisions(self, jobs: list[TTSJob]) -> list[SourceImportIssue]:
        grouped: dict[str, list[TTSJob]] = {}
        for job in jobs:
            grouped.setdefault(Path(job.filename).as_posix().casefold(), []).append(job)
        collisions: list[SourceImportIssue] = []
        for _filename, items in grouped.items():
            if len(items) < 2:
                continue
            for job in items:
                other = next((item for item in items if item is not job), items[0])
                collisions.append(
                    SourceImportIssue(
                        source_id=job.source_id or "",
                        severity="error",
                        code="cross_source_filename_collision",
                        message="Output filename collides with another enabled source.",
                        physical_row=job.source_row,
                        filename=job.filename,
                        other_source_id=other.source_id,
                    )
                )
        return collisions

    def changed_sources(self, sources: list[ProjectSource]) -> list[ProjectSource]:
        changed: list[ProjectSource] = []
        for source in sources:
            if not source.source_path.exists():
                source.import_status = SourceImportStatus.MISSING
                changed.append(source)
                continue
            current_hash = source_hash(source.source_path)
            if source.source_hash and current_hash != source.source_hash:
                source.import_status = SourceImportStatus.CHANGED
                changed.append(source)
        return changed

    def excel_worksheets(self, path: Path) -> list[str]:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise CSVValidationError("Excel import requires openpyxl. Install the 'excel' optional dependency.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def _import_delimited(self, source: ProjectSource) -> SourceImportResult:
        text, encoding = self._read_text(source.source_path)
        delimiter = "\t" if source.source_type == SourceType.TSV else self._detect_delimiter(text)
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
        source.detected_encoding = encoding
        source.detected_delimiter = delimiter
        return self._rows_to_result(source, rows)

    def _import_excel(self, source: ProjectSource) -> SourceImportResult:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise CSVValidationError("Excel import requires openpyxl. Install the 'excel' optional dependency.") from exc
        workbook = load_workbook(source.source_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[source.worksheet_name] if source.worksheet_name else workbook[workbook.sheetnames[0]]
            rows = [[self._cell_text(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
            source.worksheet_name = worksheet.title
            return self._rows_to_result(source, rows)
        finally:
            workbook.close()

    def _rows_to_result(self, source: ProjectSource, rows: list[list[str]]) -> SourceImportResult:
        result = SourceImportResult(source)
        source.imported_at = utc_now()
        if source.source_path.exists():
            source.last_modified = str(source.source_path.stat().st_mtime)
            source.source_hash = source_hash(source.source_path)
        if not rows:
            result.issues.append(self._issue(source, "error", "empty_source", "Source contains no rows."))
            self._finalize(source, result)
            return result
        headers = [self._normalize_header(value) for value in rows[0]]
        result.headers = headers
        result.preview_rows = self._preview(headers, rows[1:6])
        duplicate_headers = {name for name, count in Counter(headers).items() if name and count > 1}
        for header in sorted(duplicate_headers):
            result.issues.append(self._issue(source, "error", "duplicate_header", f"Duplicate header: {header}.", physical_row=1))
        required = [source.mapping.text_column, source.mapping.filename_column]
        missing = [column for column in required if self._normalize_header(column) not in headers]
        for column in missing:
            result.issues.append(self._issue(source, "error", "missing_column", f"Missing required column: {column}.", physical_row=1))
        if duplicate_headers or missing:
            self._finalize(source, result)
            return result
        text_index = headers.index(self._normalize_header(source.mapping.text_column))
        filename_index = headers.index(self._normalize_header(source.mapping.filename_column))
        optional_indexes = self._optional_indexes(source.mapping, headers)
        for index, row in enumerate(rows[1:], start=2):
            if source.row_end is not None and index > source.row_end:
                break
            if index < source.row_start:
                continue
            result.row_count += 1
            filename = self._get(row, filename_index)
            text = self._get(row, text_index).replace("\r\n", "\n")
            issues = self._validate_row(source, index, filename, text)
            result.issues.extend(issues)
            if any(issue.severity == "error" for issue in issues):
                continue
            job = TTSJob(
                row_number=0,
                filename=filename,
                text=text,
                source_physical_row=index,
                source_id=source.source_id,
                source_display_name=source.display_name,
                source_sheet=source.worksheet_name,
                source_row=index,
                voice_override=self._optional_value(row, optional_indexes.get("voice")),
                model_override=self._optional_value(row, optional_indexes.get("model")),
                language_override=self._optional_value(row, optional_indexes.get("language")),
                output_subfolder=self._optional_value(row, optional_indexes.get("output_subfolder")),
            )
            result.jobs.append(job)
        self._finalize(source, result)
        return result

    def _finalize(self, source: ProjectSource, result: SourceImportResult) -> None:
        source.valid_rows = len(result.jobs)
        source.rejected_rows = sum(1 for issue in result.issues if issue.severity == "error" and issue.physical_row != 1)
        source.issue_count = len(result.issues)
        if result.issues and any(issue.severity == "error" for issue in result.issues):
            source.import_status = SourceImportStatus.ERROR
        elif result.issues:
            source.import_status = SourceImportStatus.WARNING
        else:
            source.import_status = SourceImportStatus.READY

    def _source(
        self,
        path: Path,
        *,
        project_id: int | None,
        order: int,
        source_type: SourceType,
        worksheet_name: str | None = None,
    ) -> ProjectSource:
        display = path.stem if worksheet_name is None else f"{path.stem}:{worksheet_name}"
        return ProjectSource(
            source_id=uuid.uuid4().hex,
            project_id=project_id,
            display_name=display,
            source_type=source_type,
            source_path=path,
            worksheet_name=worksheet_name,
            import_order=order,
        )

    def _unsupported_source(self, path: Path, project_id: int | None, order: int, message: str) -> ProjectSource:
        source = self._source(path, project_id=project_id, order=order, source_type=SourceType.CSV)
        source.import_status = SourceImportStatus.UNSUPPORTED
        source.issue_count = 1
        return source

    @staticmethod
    def assign_merged_row_numbers(jobs: list[TTSJob]) -> list[TTSJob]:
        for index, job in enumerate(jobs, start=1):
            job.row_number = index
            job.original_order = index
        return jobs

    @staticmethod
    def _read_text(path: Path) -> tuple[str, str]:
        data = path.read_bytes()
        for encoding in ENCODINGS:
            try:
                return data.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        raise CSVValidationError("Source must be UTF-8, UTF-8 BOM, cp1252, or latin-1 encoded.")

    @staticmethod
    def _detect_delimiter(text: str) -> str:
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            return dialect.delimiter
        except csv.Error:
            return ","

    @staticmethod
    def _validate_row(source: ProjectSource, physical_row: int, filename: str, text: str) -> list[SourceImportIssue]:
        issues: list[SourceImportIssue] = []
        if not filename.strip():
            issues.append(SourceImportService._issue(source, "error", "missing_filename", "Filename is empty.", physical_row=physical_row))
        if not text.strip():
            issues.append(SourceImportService._issue(source, "error", "missing_text", "Text is empty.", physical_row=physical_row, filename=filename))
        basename = Path(filename).name
        if INVALID_WINDOWS.search(basename) or not FILENAME_EXT.search(basename):
            issues.append(SourceImportService._issue(source, "error", "invalid_filename", "Filename is invalid or has an unsupported extension.", physical_row=physical_row, filename=filename))
        return issues

    @staticmethod
    def _optional_indexes(mapping: SourceColumnMapping, headers: list[str]) -> dict[str, int | None]:
        columns = {
            "voice": mapping.voice_column,
            "model": mapping.model_column,
            "language": mapping.language_column,
            "output_subfolder": mapping.output_subfolder_column,
        }
        return {
            key: headers.index(SourceImportService._normalize_header(value))
            if value and SourceImportService._normalize_header(value) in headers
            else None
            for key, value in columns.items()
        }

    @staticmethod
    def _preview(headers: list[str], rows: list[list[str]]) -> list[dict[str, str]]:
        return [{header: SourceImportService._get(row, index) for index, header in enumerate(headers) if header} for row in rows]

    @staticmethod
    def _normalize_header(value: str | None) -> str:
        return (value or "").strip().casefold()

    @staticmethod
    def _get(row: list[str], index: int) -> str:
        return str(row[index]) if 0 <= index < len(row) and row[index] is not None else ""

    @staticmethod
    def _optional_value(row: list[str], index: int | None) -> str | None:
        if index is None:
            return None
        value = SourceImportService._get(row, index).strip()
        return value or None

    @staticmethod
    def _cell_text(value: object) -> str:
        return "" if value is None else str(value)

    @staticmethod
    def _issue(
        source: ProjectSource,
        severity: str,
        code: str,
        message: str,
        *,
        physical_row: int | None = None,
        filename: str = "",
    ) -> SourceImportIssue:
        return SourceImportIssue(source.source_id, severity, code, message, physical_row=physical_row, filename=filename)
